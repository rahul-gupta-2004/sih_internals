import subprocess
import sys
import tempfile
from pathlib import Path
import chromadb
from chromadb.utils import embedding_functions
import ollama
from pathlib import Path

BASE_DIR = Path(__file__).resolve().parent
OUTPUTS_DIR = BASE_DIR / "outputs"
SANDBOX_DIR = BASE_DIR / "sandbox"
KB_DIR = BASE_DIR / "kb_docs"
CHROMA_DIR = BASE_DIR / "chroma_db"

OUTPUTS_DIR.mkdir(exist_ok=True)
SANDBOX_DIR.mkdir(exist_ok=True)
KB_DIR.mkdir(exist_ok=True)
CHROMA_DIR.mkdir(exist_ok=True)


# ---------- Tool 1: Code Execution Sandbox ----------
def execute_python(code: str, timeout: int = 15) -> dict:
    """
    Execute Python code in an isolated subprocess with a timeout.
    Returns stdout, stderr, and exit code.
    """
    # Write code to a temp file in sandbox
    script_name = "run.py"
    script_path = SANDBOX_DIR / script_name
    script_path.write_text(code, encoding="utf-8")

    try:
        result = subprocess.run(
            [sys.executable, script_name],   # ← just the filename, not the full path
            capture_output=True,
            text=True,
            timeout=timeout,
            cwd=str(SANDBOX_DIR.resolve()),  # ← resolve() for absolute path safety
        )
        return {
            "success": result.returncode == 0,
            "stdout": result.stdout,
            "stderr": result.stderr,
            "exit_code": result.returncode,
        }
    except subprocess.TimeoutExpired:
        return {
            "success": False,
            "stdout": "",
            "stderr": f"Execution timed out after {timeout} seconds.",
            "exit_code": -1,
        }
    except Exception as e:
        return {
            "success": False,
            "stdout": "",
            "stderr": str(e),
            "exit_code": -1,
        }


# ---------- Tool 2: Knowledge Base Search ----------
_chroma_client = None
_collection = None


def _get_collection():
    """Lazy-init the ChromaDB collection with Nomic embeddings via Ollama."""
    global _chroma_client, _collection
    if _collection is not None:
        return _collection

    _chroma_client = chromadb.PersistentClient(path=str(CHROMA_DIR))

    # Custom embedding function using local Ollama Nomic model
    class OllamaEmbedding(embedding_functions.EmbeddingFunction):
        def __call__(self, input: list[str]) -> list[list[float]]:
            embeddings = []
            for text in input:
                resp = ollama.embeddings(model="nomic-embed-text:v1.5", prompt=text)
                embeddings.append(resp["embedding"])
            return embeddings

    _collection = _chroma_client.get_or_create_collection(
        name="knowledge_base",
        embedding_function=OllamaEmbedding(),
    )
    return _collection


def build_knowledge_base():
    """Index all .txt and .pdf files in kb_docs/ into ChromaDB."""
    from pypdf import PdfReader

    collection = _get_collection()
    docs = []
    ids = []
    metadatas = []

    for i, path in enumerate(sorted(KB_DIR.iterdir())):
        if path.suffix.lower() == ".txt":
            text = path.read_text(encoding="utf-8", errors="ignore")
        elif path.suffix.lower() == ".pdf":
            try:
                reader = PdfReader(str(path))
                text = "\n".join(page.extract_text() or "" for page in reader.pages)
            except Exception as e:
                text = f"[Could not read PDF: {e}]"
        else:
            continue

        # Chunk into ~500-char pieces
        chunks = [text[j:j+500] for j in range(0, len(text), 500) if text[j:j+500].strip()]
        for k, chunk in enumerate(chunks):
            docs.append(chunk)
            ids.append(f"{path.stem}_{k}")
            metadatas.append({"source": path.name})

    if not docs:
        return 0

    # Upsert (idempotent)
    collection.upsert(documents=docs, ids=ids, metadatas=metadatas)
    return len(docs)


def kb_search(query: str, top_k: int = 3) -> list[dict]:
    """Search the local knowledge base."""
    collection = _get_collection()
    if collection.count() == 0:
        return []
    results = collection.query(query_texts=[query], n_results=top_k)
    hits = []
    for doc, meta in zip(results["documents"][0], results["metadatas"][0]):
        hits.append({"text": doc, "source": meta.get("source", "unknown")})
    return hits


# ---------- Tool 3: Document Writer ----------
def write_word_doc(filename: str, title: str, body: str) -> Path:
    """Write a .docx file to outputs/."""
    from docx import Document
    doc = Document()
    doc.add_heading(title, level=1)
    for para in body.split("\n\n"):
        if para.strip():
            doc.add_paragraph(para.strip())
    out_path = OUTPUTS_DIR / filename
    doc.save(str(out_path))
    return out_path


def write_text_file(filename: str, content: str) -> Path:
    """Write a plain text file to outputs/."""
    out_path = OUTPUTS_DIR / filename
    out_path.write_text(content, encoding="utf-8")
    return out_path